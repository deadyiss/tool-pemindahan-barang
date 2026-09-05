"""
Parser untuk 2 format laporan Excel (lihat PRD §2.2):
- 'movement' : Laporan Stok - Movement Stock (SKU ditulis sebagai formula ="...")
- 'daftar'   : Daftar Stok (SKU ditulis sebagai string biasa)

Format dideteksi otomatis dari ISI file (teks judul), bukan dari nama file.
"""
import re
from dataclasses import dataclass

import openpyxl

SKU_FORMULA_PATTERN = re.compile(r'^="(.*)"$')


@dataclass
class BarangRow:
    sku: str
    nama_barang: str
    category: str | None
    qty_masuk: float      # Receive (movement) / Masuk (daftar)
    qty_terjual: float    # Penjualan (movement) / Terjual (daftar)


@dataclass
class ParsedFile:
    format_sumber: str    # 'movement' | 'daftar'
    cabang: str
    tanggal_awal: str      # ISO YYYY-MM-DD, tanggal pertama dari field Periode
    rows: list[BarangRow]


def _to_iso_date(raw: str) -> str:
    """Ambil tanggal PERTAMA dari string periode, ubah ke ISO YYYY-MM-DD.
    Menangani 2 gaya penulisan yang ditemukan di data:
      '11/08/2026 - 11/08/2026'         -> DD/MM/YYYY
      '1 August 2026 - 2 August 2026'   -> D Month YYYY
    Kalau periode berupa rentang (misal 01/08-02/08), diambil TANGGAL AWAL
    sesuai keputusan PRD §2.4.
    """
    raw = raw.strip()
    m = re.search(r"(\d{1,2})/(\d{1,2})/(\d{4})", raw)
    if m:
        d, mo, y = m.groups()
        return f"{y}-{int(mo):02d}-{int(d):02d}"

    months = {
        "january": 1, "february": 2, "march": 3, "april": 4, "may": 5, "june": 6,
        "july": 7, "august": 8, "september": 9, "october": 10, "november": 11, "december": 12,
    }
    m = re.search(r"(\d{1,2})\s+([A-Za-z]+)\s+(\d{4})", raw)
    if m:
        d, mon_name, y = m.groups()
        mo = months.get(mon_name.lower())
        if mo:
            return f"{y}-{mo:02d}-{int(d):02d}"

    raise ValueError(f"Tidak bisa parse tanggal dari periode: {raw!r}")


def detect_format(ws) -> str | None:
    """Cek 12 baris pertama x 4 kolom pertama untuk teks penanda format."""
    text = []
    for r in range(1, 13):
        for c in range(1, 5):
            v = ws.cell(row=r, column=c).value
            if v:
                text.append(str(v))
    joined = " ".join(text)
    if "Movement Stock" in joined:
        return "movement"
    if "Daftar Stok" in joined:
        return "daftar"
    return None


def _parse_movement(ws) -> ParsedFile:
    cabang = None
    periode = None
    for r in range(1, 8):
        for c in range(1, 4):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, str) and v.startswith("Cabang"):
                cabang = v.split(":", 1)[-1].strip()
            if isinstance(v, str) and v.startswith("Periode"):
                periode = v.split(":", 1)[-1].strip()
    if not cabang or not periode:
        raise ValueError("Header Cabang/Periode tidak ditemukan (format movement).")

    rows = []
    for row in ws.iter_rows(min_row=8, values_only=False):
        sku_cell = row[0].value
        nama = row[1].value
        if sku_cell is None or nama is None:
            continue
        if str(nama).strip().upper() == "TOTAL":
            continue
        m = SKU_FORMULA_PATTERN.match(str(sku_cell))
        sku = m.group(1) if m else str(sku_cell)
        if not sku.strip():
            # Baris tanpa SKU (mis. baris kategori/subtotal seperti "SLOKI") -
            # bukan barang individual yang bisa dipindahkan/dilacak per-SKU, skip.
            continue
        category = row[2].value
        stok_awal, receive, retur, transfer_in, penjualan = [
            row[i].value if row[i].value is not None else 0 for i in range(3, 8)
        ]
        rows.append(
            BarangRow(
                sku=sku,
                nama_barang=str(nama).strip(),
                category=category,
                qty_masuk=float(receive or 0),
                qty_terjual=float(penjualan or 0),
            )
        )
    return ParsedFile(
        format_sumber="movement",
        cabang=cabang,
        tanggal_awal=_to_iso_date(periode),
        rows=rows,
    )


def _parse_daftar(ws) -> ParsedFile:
    cabang = None
    periode = None
    for r in range(1, 13):
        for c in range(1, 4):
            v = ws.cell(row=r, column=c).value
            if v == "Outlet":
                cabang = ws.cell(row=r, column=c + 1).value
            if v == "Periode":
                periode = ws.cell(row=r, column=c + 1).value
    if not cabang or not periode:
        raise ValueError("Header Outlet/Periode tidak ditemukan (format daftar).")

    rows = []
    for row in ws.iter_rows(min_row=13, values_only=False):
        sku = row[1].value
        nama = row[2].value
        if sku is None or nama is None:
            continue
        masuk = row[5].value if row[5].value is not None else 0
        terjual = row[6].value if row[6].value is not None else 0
        rows.append(
            BarangRow(
                sku=str(sku),
                nama_barang=str(nama).strip(),
                category=row[3].value,
                qty_masuk=float(masuk or 0),
                qty_terjual=float(terjual or 0),
            )
        )
    return ParsedFile(
        format_sumber="daftar",
        cabang=str(cabang).strip(),
        tanggal_awal=_to_iso_date(str(periode)),
        rows=rows,
    )


def parse_file(path: str) -> ParsedFile:
    wb = openpyxl.load_workbook(path, data_only=False, read_only=False)
    ws = wb.active
    fmt = detect_format(ws)
    if fmt == "movement":
        result = _parse_movement(ws)
    elif fmt == "daftar":
        result = _parse_daftar(ws)
    else:
        raise ValueError(f"Format tidak dikenali untuk file: {path}")
    wb.close()
    return result
